# -*- coding: utf-8 -*-
"""Gera os templates Excel dos Books do BooX a partir de schema/campos.py."""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "schema"))
import campos
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

AQUI = os.path.dirname(os.path.abspath(__file__))

COR = {"OBR":"C0392B","REC":"B9770E","OPC":"5F6A6A"}  # vermelho/ambar/cinza
FILL = {k:PatternFill("solid",start_color=v,end_color=v) for k,v in COR.items()}
ROT = {"OBR":"OBRIGATORIO","REC":"RECOMENDADO","OPC":"OPCIONAL"}
THIN = Side(style="thin", color="D5D5D5")
BORDER = Border(left=THIN,right=THIN,top=THIN,bottom=THIN)

def gerar_book(nome, spec):
    wb = Workbook()
    ws = wb.active; ws.title = "Dados"
    campos_lst = spec["campos"]

    # Linha 1: categoria (merge visual nao — so cor por celula)
    # Linha 2: nome do campo (cabecalho real)
    for j,(cnome,cat,tipo,dica) in enumerate(campos_lst, start=1):
        c1 = ws.cell(row=1, column=j, value=ROT[cat])
        c1.fill = FILL[cat]; c1.font = Font(bold=True, color="FFFFFF", size=8, name="Arial")
        c1.alignment = Alignment(horizontal="center")
        c1.border = BORDER
        c2 = ws.cell(row=2, column=j, value=cnome)
        c2.font = Font(bold=True, color="FFFFFF", size=10, name="Arial")
        c2.fill = FILL[cat]; c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c2.border = BORDER
        # comentario com a dica
        dica_txt = dica if tipo!="lista" else f"{dica}\nValores: " + ", ".join(campos.ENUM[dica])
        c2.comment = Comment(dica_txt, "BooX")
        # largura
        ws.column_dimensions[c2.column_letter].width = max(14, min(len(cnome)+4, 28))
        # validacao por lista
        if tipo=="lista":
            opts = campos.ENUM[dica]
            dv = DataValidation(type="list", formula1='"'+",".join(opts)+'"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{c2.column_letter}3:{c2.column_letter}1000")
        elif tipo=="bool":
            dv = DataValidation(type="list", formula1='"true,false"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f"{c2.column_letter}3:{c2.column_letter}1000")
    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 14; ws.row_dimensions[2].height = 30

    # aba de instrucoes
    wi = wb.create_sheet("Instrucoes")
    wi.column_dimensions["A"].width = 24; wi.column_dimensions["B"].width = 14
    wi.column_dimensions["C"].width = 16; wi.column_dimensions["D"].width = 70
    wi["A1"] = spec["titulo"]; wi["A1"].font = Font(bold=True, size=14, name="Arial")
    wi["A2"] = spec["desc"]; wi["A2"].font = Font(italic=True, size=10, name="Arial", color="555555")
    hdr = ["Campo","Categoria","Tipo","Descricao / valores aceitos"]
    for j,h in enumerate(hdr, start=1):
        c = wi.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF", name="Arial"); c.fill = PatternFill("solid",start_color="34495E",end_color="34495E")
        c.border = BORDER
    r = 5
    for cnome,cat,tipo,dica in campos_lst:
        wi.cell(row=r,column=1,value=cnome).font=Font(name="Arial",bold=(cat=="OBR"))
        cc=wi.cell(row=r,column=2,value=ROT[cat]); cc.font=Font(name="Arial",color="FFFFFF",size=9,bold=True); cc.fill=FILL[cat]; cc.alignment=Alignment(horizontal="center")
        wi.cell(row=r,column=3,value=tipo).font=Font(name="Arial")
        d = dica if tipo!="lista" else dica+" -> " + ", ".join(campos.ENUM[dica])
        wi.cell(row=r,column=4,value=d).font=Font(name="Arial",size=9)
        for col in range(1,5): wi.cell(row=r,column=col).border=BORDER
        r+=1
    # legenda
    r+=1
    wi.cell(row=r,column=1,value="LEGENDA").font=Font(bold=True,name="Arial"); r+=1
    leg=[("OBRIGATORIO","Sem este campo o registro NAO entra no 20-box","OBR"),
         ("RECOMENDADO","A ausencia reduz a qualidade/confianca, mas nao impede o calculo","REC"),
         ("OPCIONAL","Enriquece a analise; nao afeta o calculo","OPC")]
    for nome_l,desc_l,cat in leg:
        c=wi.cell(row=r,column=1,value=nome_l); c.fill=FILL[cat]; c.font=Font(color="FFFFFF",bold=True,name="Arial",size=9); c.alignment=Alignment(horizontal="center")
        wi.cell(row=r,column=2,value=desc_l).font=Font(name="Arial",size=9); r+=1

    caminho = os.path.join(AQUI, f"{nome}_template.xlsx")
    wb.save(caminho)
    return caminho

# Agrupar os arquivos B6 e B10 e B4 e B7 que tem 2 tabelas em 1 workbook cada? 
# Mantemos 1 arquivo por tabela para simplicidade do preenchimento.
gerados = []
for nome, spec in campos.BOOKS.items():
    gerados.append(gerar_book(nome, spec))
print(f"{len(gerados)} templates gerados:")
for g in gerados: print("  ", os.path.basename(g))
